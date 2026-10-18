"""Build the Assignment 01 report as a Word document (report.docx).

Usage (from the repo root, venv active):
    pip install python-docx reportlab
    python scripts/build_report_docx.py            # builds report.docx
    python scripts/build_report_docx.py --strict   # also exits with an error if anything is unfinished

It uses the SAME content as scripts/build_report.py: your name, ID, links and the USER_TEXT block
are read from that file, so you only edit them in one place. Everything still unfinished is
highlighted in yellow, and the terminal lists it.

After opening report.docx in Word: press Ctrl+A, then F9 (Update fields) once, so the contents
page gets its page numbers.
"""
import argparse
import html
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_report as br  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402
from reportlab.platypus import (Image, KeepTogether, PageBreak, Paragraph, Preformatted,  # noqa: E402
                                Spacer, Table)
from reportlab.platypus.tableofcontents import TableOfContents  # noqa: E402

OUT_DOCX = br.ROOT / "report.docx"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GRID = "B0B7C3"
TAG = re.compile(r"(<[^>]+>)")


# ---------------------------------------------------------------------------------
#  Small helpers
# ---------------------------------------------------------------------------------
def add_rich(par, text, size=None, bold=False, italic=False, color=None, font=None):
    """Add text containing the few markup tags used in the report (<b>, <i>, <font backColor>)."""
    b, i, hl = bold, italic, 0
    for tok in TAG.split(text):
        if not tok:
            continue
        if tok.startswith("<"):
            t = tok.lower()
            if t == "<b>":
                b = True
            elif t == "</b>":
                b = bold
            elif t == "<i>":
                i = True
            elif t == "</i>":
                i = italic
            elif t.startswith("<font") and "backcolor" in t:
                hl += 1
            elif t == "</font>":
                hl = max(0, hl - 1)
            elif t.startswith("<br"):
                par.add_run().add_break()
            continue
        run = par.add_run(html.unescape(tok))
        run.bold = b or None
        run.italic = i or None
        if size:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
        if font:
            run.font.name = font
        if hl:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def shade(element_pr, fill_hex):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    element_pr.append(shd)


def para_shade(par, fill_hex):
    shade(par._p.get_or_add_pPr(), fill_hex)


def set_borders(table, kind):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    names = ["top", "left", "bottom", "right"] + (["insideH", "insideV"] if kind == "grid" else [])
    for n in names:
        el = OxmlElement(f"w:{n}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), GRID)
        borders.append(el)
    anchor = None
    for tag in ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook"):
        anchor = tbl_pr.find(qn(tag))
        if anchor is not None:
            break
    if anchor is not None:
        anchor.addprevious(borders)
    else:
        tbl_pr.append(borders)


def field_run(par, instr_text, placeholder="1"):
    """Insert a Word field (PAGE etc.) into a paragraph."""
    def fld(kind):
        r = par.add_run()
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        r._r.append(el)
        return r
    fld("begin")
    r = par.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr_text} "
    r._r.append(it)
    fld("separate")
    par.add_run(placeholder)
    fld("end")


def hexcolor(c):
    return c.hexval()[2:].upper()


# ---------------------------------------------------------------------------------
#  Converter: reportlab flowables -> Word
# ---------------------------------------------------------------------------------
class Converter:
    def __init__(self, doc, h1_titles):
        self.doc = doc
        self.h1_titles = h1_titles

    # -- paragraphs ---------------------------------------------------------------
    def paragraph(self, f, container=None, first_in_cell=None):
        container = container or self.doc
        name = f.style.name
        if first_in_cell is not None:
            par = first_in_cell
        elif name == "H1x":
            par = container.add_paragraph(style="Heading 1")
        elif name == "H2x":
            par = container.add_paragraph(style="Heading 2")
        else:
            par = container.add_paragraph()
        size = None if name in ("H1x", "H2x") else f.style.fontSize
        bold = "Bold" in f.style.fontName
        color = None
        tc = f.style.textColor
        if name not in ("H1x", "H2x") and (tc.red, tc.green, tc.blue) != (0, 0, 0):
            color = RGBColor(int(tc.red * 255), int(tc.green * 255), int(tc.blue * 255))
        if name in ("TitleX", "Sub", "TocTitle"):
            color = NAVY if name != "Sub" else None
        add_rich(par, f.text, size=size, bold=bold or name == "TitleX", color=color,
                 italic=(name == "Caption"))
        pf = par.paragraph_format
        if name not in ("H1x", "H2x"):
            pf.space_before = Pt(f.style.spaceBefore or 0)
            pf.space_after = Pt(f.style.spaceAfter if f.style.spaceAfter is not None else 6)
            if f.style.alignment == 1:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if name == "TocTitle":
            pf.keep_with_next = True
        return par

    def code(self, f):
        par = self.doc.add_paragraph()
        run = par.add_run()
        run.font.name = "Consolas"
        run.font.size = Pt(7.8)
        for n, line in enumerate(f.lines):
            if n:
                run.add_break()
            run.add_text(line)
        para_shade(par, "F3F3F3")
        par.paragraph_format.space_after = Pt(8)
        par.paragraph_format.keep_together = True
        return par

    def image(self, f, container=None, par=None):
        container = container or self.doc
        par = par or container.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.add_run().add_picture(f.filename, width=Pt(f.drawWidth))
        par.paragraph_format.space_after = Pt(2)
        return par

    # -- tables ---------------------------------------------------------------------
    def table(self, t):
        rows, cols = len(t._cellvalues), len(t._cellvalues[0])
        tb = self.doc.add_table(rows=rows, cols=cols)
        tb.autofit = False
        kinds = {c[0] for c in t._linecmds}
        if "GRID" in kinds:
            set_borders(tb, "grid")
        elif "BOX" in kinds:
            set_borders(tb, "box")
        widths = t._colWidths
        for r in range(rows):
            trpr = tb.rows[r]._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            for c in range(cols):
                cell = tb.cell(r, c)
                if widths and widths[c]:
                    cell.width = Pt(widths[c])
                self.cell_content(cell, t._cellvalues[r][c])
        for cmd in t._bkgrndcmds:
            (c0, r0), (c1, r1) = cmd[1], cmd[2]
            c0, c1 = c0 % cols, c1 % cols
            r0, r1 = r0 % rows, r1 % rows
            for r in range(r0, r1 + 1):
                for c in range(c0, c1 + 1):
                    shade(tb.cell(r, c)._tc.get_or_add_tcPr(), hexcolor(cmd[3]))
        # small gap after the table
        gap = self.doc.add_paragraph()
        gap.paragraph_format.space_after = Pt(2)
        for run in gap.runs:
            run.font.size = Pt(2)
        return tb

    def cell_content(self, cell, value):
        items = value if isinstance(value, list) else [value]
        first = True
        for it in items:
            if isinstance(it, str):
                it = Paragraph(html.escape(it), br.CELL)
            if isinstance(it, Spacer):
                continue
            par = cell.paragraphs[0] if first else None
            if isinstance(it, Paragraph):
                p = self.paragraph(it, container=cell, first_in_cell=par) if first \
                    else self.paragraph(it, container=cell)
                p.paragraph_format.space_after = Pt(it.style.spaceAfter or 2)
                if p.paragraph_format.space_before is None:
                    p.paragraph_format.space_before = Pt(0)
            elif isinstance(it, Image):
                p = self.image(it, container=cell, par=par)
            else:
                continue
            first = False

    # -- main walk ------------------------------------------------------------------
    def walk(self, flowables):
        made = []
        for f in flowables:
            if isinstance(f, KeepTogether):
                inner = self.walk(f._content)
                for p in inner[:-1]:
                    p.paragraph_format.keep_with_next = True
                made += inner
            elif isinstance(f, TableOfContents):
                made += self.toc()
            elif isinstance(f, Paragraph):
                made.append(self.paragraph(f))
            elif isinstance(f, Preformatted):
                made.append(self.code(f))
            elif isinstance(f, Image):
                made.append(self.image(f))
            elif isinstance(f, Table):
                self.table(f)
            elif isinstance(f, PageBreak):
                self.doc.add_page_break()
            elif isinstance(f, Spacer):
                if f.height >= 8:
                    p = self.doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = Pt(f.height)
        return made

    def toc(self):
        """A real Word contents field, pre-filled with the section titles (F9 adds page numbers)."""
        made = []
        for n, title in enumerate(self.h1_titles):
            par = self.doc.add_paragraph()
            par.paragraph_format.space_after = Pt(3)
            if n == 0:
                def fld(kind):
                    r = par.add_run()
                    el = OxmlElement("w:fldChar")
                    el.set(qn("w:fldCharType"), kind)
                    r._r.append(el)
                fld("begin")
                r = par.add_run()
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = ' TOC \\o "1-1" \\h \\z \\u '
                r._r.append(it)
                fld("separate")
            par.add_run(title)
            if n == len(self.h1_titles) - 1:
                r = par.add_run()
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), "end")
                r._r.append(el)
            made.append(par)
        return made


def flatten(flowables):
    for f in flowables:
        if isinstance(f, KeepTogether):
            yield from flatten(f._content)
        else:
            yield f


def setup_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2)
    sec.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for style_name, size, before, after in (("Heading 1", 15, 14, 8), ("Heading 2", 11.5, 8, 4)):
        st = doc.styles[style_name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = NAVY
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    foot = sec.footer.paragraphs[0]
    run = foot.add_run("MLOps Assignment 01 - Heart Disease Risk Prediction        Page ")
    run.font.size = Pt(8)
    field_run(foot, "PAGE")
    for r in foot.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    zoom = doc.settings.element.find(qn("w:zoom"))      # python-docx template omits this attribute
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    doc.core_properties.title = "MLOps Assignment 01 - Heart Disease Risk Prediction"
    doc.core_properties.author = ""
    return doc


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--strict", action="store_true", help="exit with an error if anything is unfinished")
    args = ap.parse_args()

    br.SHOTS.mkdir(exist_ok=True)
    print("Drawing architecture diagram ...")
    br.draw_architecture(br.SHOTS / "architecture.png")
    print("Reading repository data ...")
    stats, runs, meta = br.data_stats(), br.load_mlflow(), br.load_metadata()
    print("Building Word document ...")
    story = br.build_story(stats, runs, meta)
    titles = [f.getPlainText() for f in flatten(story)
              if isinstance(f, Paragraph) and f.style.name == "H1x"]

    doc = setup_document()
    Converter(doc, titles).walk(story)
    doc.save(str(OUT_DOCX))
    print(f"\nWrote {OUT_DOCX.name}")
    print("Open it in Word, press Ctrl+A then F9 once to update the contents page numbers.")

    if br.MISSING_SHOTS:
        print(f"\nMissing screenshots ({len(br.MISSING_SHOTS)}):")
        for m in dict.fromkeys(br.MISSING_SHOTS):
            print("   screenshots/" + m)
    if br.PLACEHOLDERS:
        print(f"\nStill to write ({len(br.PLACEHOLDERS)} highlighted placeholders):")
        for p in dict.fromkeys(br.PLACEHOLDERS):
            print("   -", p[:90] + ("..." if len(p) > 90 else ""))
    if not br.MISSING_SHOTS and not br.PLACEHOLDERS:
        print("\nNothing left to fill in. Review the document once, start to finish.")
    if args.strict and (br.MISSING_SHOTS or br.PLACEHOLDERS):
        sys.exit(1)


if __name__ == "__main__":
    main()
